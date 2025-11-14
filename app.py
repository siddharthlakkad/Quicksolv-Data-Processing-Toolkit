import streamlit as st
import pandas as pd
import openpyxl
from io import BytesIO

# ------------------------
# APP CONFIG
# ------------------------
st.set_page_config(page_title="Data Tools", layout="wide")

st.markdown(
    """
    <style>
    .module-box {
        padding: 20px;
        background: #ffffff;
        border-radius: 12px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.08);
        margin-bottom: 20px;
    }
    .title-main {
        font-size: 30px; 
        text-align: center; 
        font-weight: 700; 
        color: #1a73e8;
    }
    </style>
    """,
    unsafe_allow_html=True
)

st.markdown("<div class='title-main'>📌 Data Processing Toolkit</div>", unsafe_allow_html=True)

# ------------------------
# SIDEBAR MENU
# ------------------------
tool = st.sidebar.radio(
    "Choose Tool",
    ["Excel Cleaner (Cross-File)", "Straight-Line → Comma Converter"]
)


# ================================================================
# 1️⃣ MODULE : Excel Cleaner (Cross-File)
# ================================================================
def module_excel_cleaner():
    st.markdown("<div class='module-box'>", unsafe_allow_html=True)
    st.subheader("🧹 Excel Cleaner — Remove rows from Main File based on Data File")

    main_file = st.file_uploader("Upload Main Excel File (.xlsx)", type=["xlsx"], key="main_file")
    data_file = st.file_uploader("Upload Data File (.xlsx / .csv / .txt)", type=["xlsx", "csv", "txt"], key="data_file")

    if main_file and data_file:

        # --- Load Main File Sheets ---
        try:
            main_sheets = openpyxl.load_workbook(main_file, read_only=True).sheetnames
        except:
            main_sheets = [0]
        main_sheet = st.selectbox("Select Sheet from Main File", main_sheets)

        # --- Load Main DataFrame ---
        df_main = pd.read_excel(main_file, sheet_name=main_sheet)

        # --- Load Data File (txt/csv/xlsx) ---
        data_name = data_file.name.lower()

        if data_name.endswith(".txt"):
            raw_text = data_file.getvalue().decode("utf-8")
            df_data = pd.DataFrame({"value": [line.strip() for line in raw_text.splitlines() if line.strip()]})
        elif data_name.endswith(".csv"):
            df_data = pd.read_csv(data_file)
        else:
            sheets_data = openpyxl.load_workbook(data_file, read_only=True).sheetnames
            data_sheet = st.selectbox("Select Sheet from Data File", sheets_data)
            df_data = pd.read_excel(data_file, sheet_name=data_sheet)

        st.write("### Main File Preview:")
        st.dataframe(df_main, use_container_width=True)

        st.write("### Data File Preview:")
        st.dataframe(df_data, use_container_width=True)

        # --- Column Selection ---
        main_col = st.selectbox("Select Column from Main File", df_main.columns)
        data_col = st.selectbox("Select Column from Data File", df_data.columns)

        case_insensitive = st.checkbox("Case-insensitive match", True)
        dedupe = st.checkbox("Remove duplicate values from Data File before matching", True)

        if st.button("🚀 Clean Main File", use_container_width=True):

            main_series = df_main[main_col].astype(str)
            data_series = df_data[data_col].astype(str)

            if case_insensitive:
                main_series = main_series.str.strip().str.lower()
                data_series = data_series.str.strip().str.lower()
            else:
                main_series = main_series.str.strip()
                data_series = data_series.str.strip()

            if dedupe:
                data_values = set(data_series.unique())
            else:
                data_values = data_series.tolist()

            before = len(df_main)
            mask = ~main_series.isin(data_values)
            df_cleaned = df_main[mask]
            removed = before - len(df_cleaned)

            st.success(f"✅ Removed {removed} matching rows!")

            st.dataframe(df_cleaned, use_container_width=True)

            # --- Download cleaned file ---
            buffer = BytesIO()
            with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
                df_cleaned.to_excel(writer, index=False, sheet_name="cleaned")

            st.download_button(
                "📥 Download Cleaned File",
                data=buffer.getvalue(),
                file_name="cleaned_main_file.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
    else:
        st.info("Upload both Main and Data files to continue.")

    st.markdown("</div>", unsafe_allow_html=True)



# ================================================================
# 2️⃣ MODULE : Straight-Line → Comma Converter
# ================================================================
def module_line_to_comma():
    st.markdown("<div class='module-box'>", unsafe_allow_html=True)
    st.subheader("↕️ Straight-Line → Comma Converter")

    text_input = st.text_area("Paste line-by-line text:", height=250)

    remove_dup = st.checkbox("Remove duplicates", True)
    trim_spaces = st.checkbox("Trim spaces", True)

    if st.button("Convert to comma-separated list", use_container_width=True):
        if not text_input.strip():
            st.warning("Please paste some text first!")
        else:
            items = [x.strip() for x in text_input.splitlines() if x.strip()]

            if trim_spaces:
                items = [x.strip() for x in items]

            if remove_dup:
                items = list(dict.fromkeys(items))

            result = ", ".join(items)

            st.success("Conversion Successful!")
            st.code(result)

            st.download_button(
                "⬇ Download Result (TXT)",
                result.encode("utf-8"),
                file_name="comma_output.txt",
                mime="text/plain"
            )

    st.markdown("</div>", unsafe_allow_html=True)



# ================================================================
# MAIN RENDER
# ================================================================
if tool == "Excel Cleaner (Cross-File)":
    module_excel_cleaner()

elif tool == "Straight-Line → Comma Converter":
    module_line_to_comma()
